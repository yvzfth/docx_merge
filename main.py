"""Aggregate Summary sections from DOCX files while preserving formatting.

This script extracts only the "Summary" section from each .docx under a parent
folder (searching subfolders) and merges them into a single document. The merge
retains original styles, fonts, and images by trimming in-place and composing
documents rather than recreating content.
"""

from __future__ import annotations

import argparse
import sys
import re
from io import BytesIO
from pathlib import Path
from typing import List, Optional, Sequence, Tuple, Union, Callable

from docx import Document
from docxcompose.composer import Composer
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.table import _Cell as TableCell  # type: ignore
from docx.text.paragraph import Paragraph

BlockItem = Union[Paragraph, Table]
P_TAG = qn("w:p")
TBL_TAG = qn("w:tbl")
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _clear_document(doc: Document) -> None:
    """Remove every block-level element from a document."""

    body = doc.element.body
    for child in list(body.iterchildren()):
        body.remove(child)


def _iter_block_items(doc: Document) -> List[BlockItem]:
    """Return block-level elements (paragraphs and tables) in order."""

    body = doc.element.body
    blocks: List[BlockItem] = []
    for child in body.iterchildren():
        if child.tag == P_TAG:
            blocks.append(Paragraph(child, doc))
        elif child.tag == TBL_TAG:
            blocks.append(Table(child, doc))
    return blocks


def _para_is_heading_like(p: Paragraph) -> bool:
    """Heuristic for whether a paragraph is a section heading.

    Uses style name (Heading X) or presence of w:outlineLvl.
    """

    try:
        style_name = p.style.name if p.style is not None else ""
    except Exception:
        style_name = ""
    style_lower = style_name.lower()
    if style_lower.startswith("heading"):
        return True
    # Check outline level in XML
    try:
        has_outline = bool(p._element.xpath('./w:pPr/w:outlineLvl', namespaces={"w": W_NS}))  # type: ignore[attr-defined]
        if has_outline:
            return True
    except Exception:
        pass
    return False


def _paragraph_has_summary(p: Paragraph) -> Tuple[bool, bool]:
    """Return (has_summary_text, is_heading_like) for a paragraph."""

    text = p.text.strip()
    if not text:
        return False, False
    lower = text.lower()
    if "summary" not in lower:
        return False, False
    # For anchor preference, treat true headings and styles containing "Title" as heading-like
    is_heading_like = _para_is_heading_like(p)
    try:
        style_name = p.style.name if p.style is not None else ""
        if "title" in style_name.lower():
            is_heading_like = True
    except Exception:
        pass
    return True, is_heading_like


def _paragraph_heading_level(p: Paragraph) -> Optional[int]:
    """Return heading level (0-based) if paragraph is a heading, else None.

    Uses w:outlineLvl when present; otherwise infers from style name like
    "Heading 1", "Heading 2", or treats "Title" as level 0.
    """

    # Prefer explicit outline level
    try:
        nodes = p._element.xpath('./w:pPr/w:outlineLvl', namespaces={"w": W_NS})  # type: ignore[attr-defined]
        if nodes:
            # outlineLvl@val: 0 => Heading 1, 1 => Heading 2, ...
            node = nodes[0]
            val = node.get(qn('w:val'))  # type: ignore[arg-type]
            if val is not None and val.isdigit():
                return int(val)
    except Exception:
        pass

    # Fallback: parse style name
    try:
        style_name = p.style.name if p.style is not None else ""
    except Exception:
        style_name = ""
    lower = style_name.lower()
    if lower.startswith("heading"):
        m = re.search(r"(\d+)$", style_name.strip())
        if m:
            # Heading 1 => level 0
            return max(0, int(m.group(1)) - 1)
        # Generic "Heading" without number: treat as level 0
        return 0
    if "title" in lower:
        return 0
    return None


def _text_is_headingish(text: str) -> bool:
    """Heuristic for whether plain text looks like a section heading."""

    stripped = text.strip()
    if not stripped:
        return False
    if len(stripped) > 90:
        return False
    words = stripped.split()
    if len(words) > 15:
        return False
    terminal = stripped[-1]
    if terminal in ".!?;":
        return False
    punctuation_hits = sum(1 for ch in stripped if ch in ".!?;")
    if punctuation_hits >= 2:
        return False
    return True


def _paragraph_font_size_hint(p: Paragraph) -> Optional[int]:
    """Return an effective font size hint in half-points for a paragraph.

    Priority order (max taken across):
    1) Explicit run sizes (p.runs[*].font.size)
    2) Paragraph default run properties (XML w:pPr/w:rPr/w:sz)
    3) Run rPr sizes via XML (w:r/w:rPr/w:sz)
    4) Paragraph style's font.size, following base_style chain
    """

    def _len_to_half_points(val: object) -> Optional[int]:
        # Convert python-docx Length values or numeric-like to half-points
        try:
            # python-docx Length has .pt property returning float points
            pt = getattr(val, "pt", None)
            if pt is not None:
                return int(round(float(pt) * 2))
            # Sometimes it's already an int representing EMUs or half-points; best-effort
            if isinstance(val, (int, float)):
                # Assume points if small, else assume already half-points
                if float(val) < 100:  # likely points
                    return int(round(float(val) * 2))
                return int(val)
        except Exception:
            return None
        return None

    def _style_chain_size_half_points(style_obj) -> Optional[int]:
        seen = set()
        cur = style_obj
        while cur is not None and id(cur) not in seen:
            seen.add(id(cur))
            try:
                sz = getattr(getattr(cur, "font", None), "size", None)
                hp = _len_to_half_points(sz)
                if hp is not None:
                    return hp
                cur = getattr(cur, "base_style", None)
            except Exception:
                break
        return None

    try:
        sizes: List[int] = []

        # 1) Explicit run sizes via python-docx API
        try:
            for run in getattr(p, "runs", []) or []:
                hp = _len_to_half_points(getattr(getattr(run, "font", None), "size", None))
                if hp is not None:
                    sizes.append(hp)
        except Exception:
            pass

        # 2) Paragraph default run properties size (XML)
        for node in p._element.xpath('./w:pPr/w:rPr/w:sz', namespaces={"w": W_NS}):  # type: ignore[attr-defined]
            val = node.get(qn('w:val'))  # type: ignore[arg-type]
            if val is not None and val.isdigit():
                sizes.append(int(val))

        # 3) Individual run sizes (XML)
        for node in p._element.xpath('.//w:r/w:rPr/w:sz', namespaces={"w": W_NS}):  # type: ignore[attr-defined]
            val = node.get(qn('w:val'))  # type: ignore[arg-type]
            if val is not None and val.isdigit():
                sizes.append(int(val))

        # 4) Style chain size
        try:
            style_hp = _style_chain_size_half_points(getattr(p, "style", None))
            if style_hp is not None:
                sizes.append(style_hp)
        except Exception:
            pass

        if sizes:
            return max(sizes)
    except Exception:
        pass
    return None


def _element_text_contains_summary(el: OxmlElement) -> bool:
    """Deep-scan an XML element for any text node containing 'summary' (case-insensitive).

    Captures text inside text boxes (w:txbxContent), SDTs, and nested structures.
    """

    try:
        for t in el.xpath('.//w:t', namespaces={"w": W_NS}):  # type: ignore[attr-defined]
            text = (t.text or '').strip().lower()
            if 'summary' in text:
                return True
    except Exception:
        return False
    return False


def _block_contains_summary_deep(block: BlockItem) -> Tuple[bool, bool]:
    """Return (has_summary_text_anywhere, is_heading_like) for a block.

    - For Paragraph: check paragraph text first, then deep scan descendants.
    - For Table: deep scan the XML; heading-likeness defaults to False.
    """

    if isinstance(block, Paragraph):
        found, is_heading_like = _paragraph_has_summary(block)
        if found:
            return True, is_heading_like
        return _element_text_contains_summary(block._element), False  # type: ignore[attr-defined]
    if isinstance(block, Table):
        return _element_text_contains_summary(block._element), False  # type: ignore[attr-defined]
    return False, False


def _table_has_heading_like_ge(tbl: Table, anchor_level: Optional[int], anchor_font_size: Optional[int]) -> bool:
    """Heuristically detect if a table starts a new section by containing a heading.

    Checks the first few paragraphs across cells. Returns True if it finds a
    paragraph that looks like a heading with level <= anchor_level or with
    font size >= anchor heading's size.
    """

    try:
        examined = 0
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text.strip()
                    if not text:
                        continue
                    lower = text.lower()
                    if "summary" in lower:
                        continue
                    headingish = _para_is_heading_like(p) or _text_is_headingish(text)
                    lvl = _paragraph_heading_level(p)
                    if lvl is not None and anchor_level is not None and lvl <= anchor_level:
                        return True
                    if headingish and anchor_font_size is not None:
                        size = _paragraph_font_size_hint(p)
                        if size is not None and size >= anchor_font_size:
                            return True
                    examined += 1
                    if examined >= 5:
                        return False
    except Exception:
        return False
    return False


def _select_summary_anchor(blocks: Sequence[BlockItem]) -> Optional[Tuple[int, bool]]:
    """Find the most likely block index containing the Summary heading.

    Returns (index, is_table). Heuristic prefers heading-styled content and
    blocks near the top/bottom.
    """

    hits: List[Tuple[int, int, int, bool]] = []
    total = len(blocks)
    if total == 0:
        return None

    for idx, block in enumerate(blocks):
        found, heading_like = _block_contains_summary_deep(block)
        is_table = isinstance(block, Table)
        if not found:
            continue
        heading_priority = 0 if heading_like else 1
        edge_distance = min(idx, total - idx - 1)
        hits.append((heading_priority, edge_distance, idx, is_table))

    if not hits:
        return None

    hits.sort()
    _, _, idx, is_table = hits[0]
    return idx, is_table


def _locate_summary_range(blocks: Sequence[BlockItem]) -> Optional[Tuple[int, int]]:
    """Return (start, end) indices for the Summary section inclusive of heading.

    We scan forward from the located Summary heading until the next heading
    (exclusive) or the document end.
    """

    anchor_info = _select_summary_anchor(blocks)
    if anchor_info is None:
        return None
    anchor, anchor_is_table = anchor_info

    # Determine the anchor heading level and size when the anchor is a paragraph.
    anchor_level: Optional[int] = None
    anchor_font_size: Optional[int] = None
    anchor_style_id: Optional[str] = None
    anchor_style_name: Optional[str] = None
    if isinstance(blocks[anchor], Paragraph):
        anchor_level = _paragraph_heading_level(blocks[anchor])
        anchor_font_size = _paragraph_font_size_hint(blocks[anchor])
        try:
            style = blocks[anchor].style
            anchor_style_id = getattr(style, "style_id", None)
            anchor_style_name = getattr(style, "name", None)
        except Exception:
            pass
    # Fallback: if we can't determine a level (e.g., custom style), treat Summary
    # as a mid-level heading (Heading 2 => level 1) so nested subheadings are kept.
    if anchor_level is None:
        anchor_level = 1

    end_idx = anchor + 1
    total = len(blocks)

    stopped_due_to_heading = False

    for idx in range(anchor + 1, total):
        block = blocks[idx]
        if isinstance(block, Paragraph):
            text = block.text.strip()
            lower = text.lower()
            
            # Skip empty paragraphs - include them but don't treat as boundaries
            if not text:
                end_idx = idx + 1
                continue
            
            # Check if this paragraph is a heading (multiple criteria)
            is_structural_heading = _para_is_heading_like(block)
            next_level = _paragraph_heading_level(block)
            
            # Get block style information
            try:
                block_style = block.style
                block_style_id = getattr(block_style, "style_id", None)
                block_style_name = getattr(block_style, "name", None)
            except Exception:
                block_style_id = None
                block_style_name = None
            
            # Check font size
            next_size = _paragraph_font_size_hint(block) if anchor_font_size is not None else None
            
            # Determine if this is a peer section header (same or higher level)
            is_peer_header = False
            
            # PRIORITY 1: Same style as Summary = peer header (unless it's just content with same style)
            # Check style ID independently
            if anchor_style_id is not None and block_style_id == anchor_style_id:
                # Same style: only treat as header if it's also structurally a heading
                # OR if it has heading-like text characteristics
                if is_structural_heading or _text_is_headingish(text):
                    is_peer_header = True
            # Check style name independently (may match even if ID doesn't)
            if anchor_style_name is not None and block_style_name == anchor_style_name:
                if is_structural_heading or _text_is_headingish(text):
                    is_peer_header = True
            
            # PRIORITY 2: Structural heading level check
            if next_level is not None and next_level <= anchor_level:
                is_peer_header = True
            
            # PRIORITY 3: Font size match + heading-like appearance
            if not is_peer_header and next_level is None and anchor_font_size is not None:
                if next_size is not None and next_size >= anchor_font_size:
                    # Only treat as header if text looks heading-like
                    if _text_is_headingish(text) or is_structural_heading:
                        is_peer_header = True
            
            # Stop if we found a peer header (and it's not just mentioning "summary")
            if is_peer_header and "summary" not in lower:
                stopped_due_to_heading = True
                break
            
            # Otherwise, include this paragraph in the summary section
            end_idx = idx + 1
            
        elif isinstance(block, Table):
            # Tables are included unless they contain a clear section header
            if _table_has_heading_like_ge(block, anchor_level, anchor_font_size):
                stopped_due_to_heading = True
                break
            end_idx = idx + 1

    # Fallback: if we only captured the anchor (no following blocks), widen the range
    # to include several subsequent blocks to avoid returning just the heading.
    if end_idx == anchor + 1 and not stopped_due_to_heading:
        widen_to = min(total, anchor + 6)  # keep up to 5 blocks after the anchor
        end_idx = max(end_idx, widen_to)

    return anchor, end_idx


def _remove_blocks_outside_range(doc: Document, blocks: Sequence[BlockItem], start: int, end: int) -> None:
    """Remove all block-level elements from ``doc`` that are outside [start, end).

    Removing instead of copying preserves relationships (images, styles) because
    we keep the original XML and only drop out-of-scope parts.
    """

    # Remove from the end to avoid index shifting
    for idx in range(len(blocks) - 1, -1, -1):
        if idx < start or idx >= end:
            element = blocks[idx]._element  # type: ignore[attr-defined]
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)

def _trim_document_to_range_in_place(
    doc: Document, blocks: Sequence[BlockItem], start: int, end: int
) -> Document:
    """Keep only the specified [start, end) block range in-place on ``doc``.

    This preserves all relationships, images, and formatting because content
    is removed at the XML level rather than recreated.
    """

    if start >= end:
        raise ValueError("Invalid summary range")

    _remove_blocks_outside_range(doc, blocks, start, end)
    return doc


def _label_summary_heading_with_filename(doc: Document, filename: str) -> None:
    """Append the source filename in parentheses to the Summary heading.

    This avoids resetting paragraph runs to preserve formatting. If the label
    already contains the filename, it is left unchanged.
    """

    # Prefer a heading-styled paragraph containing 'summary'
    target_para: Optional[Paragraph] = None
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if "summary" not in text.lower():
            continue
        style_name = p.style.name if p.style is not None else ""
        if style_name.lower().startswith("heading"):
            target_para = p
            break

    # Fallback: any paragraph containing 'summary'
    if target_para is None:
        for p in doc.paragraphs:
            if "summary" in p.text.lower():
                target_para = p
                break

    if target_para is None:
        return

    # Remove .docx extension from filename for display
    display_filename = filename
    if display_filename.lower().endswith('.docx'):
        display_filename = display_filename[:-5]  # Remove '.docx' (5 characters)

    suffix = f" ({display_filename})"
    if target_para.text.endswith(suffix) or f"({display_filename})" in target_para.text:
        return

    target_para.add_run(suffix)


def _build_title_doc(label: str) -> Document:
    """Create a minimal document with a heading for the given label."""

    doc = Document()
    _clear_document(doc)
    doc.add_paragraph(label, style="Heading 1")
    return doc


def _relative_label(doc_path: Path, root: Path) -> str:
    """Build a human-friendly label for the summary block."""

    try:
        rel_path = doc_path.relative_to(root)
    except ValueError:
        rel_path = doc_path.name
    return f"Summary · {rel_path.as_posix()}"


def _natural_sort_key(s: str) -> List[Union[str, int]]:
    """Return a list key that compares numbers numerically for natural sorting.

    Example: document2 < document10 (since 2 < 10), not lexicographic order.
    """

    parts = re.split(r"(\d+)", s)
    key: List[Union[str, int]] = []
    for part in parts:
        if part.isdigit():
            key.append(int(part))
        else:
            key.append(part)
    return key


def _gather_docx_files(root: Path, include_root_files: bool, output_path: Path) -> List[Path]:
    """Collect all candidate DOCX files beneath the root directory.

    Filters out temporary files (like "~$...") and the output file itself to
    avoid self-inclusion. Optionally skips .docx files in the root folder if
    only subfolders should be considered.
    """

    pattern = "**/*.docx"
    candidates: List[Path] = []
    # Natural sort by relative path string so document10 comes after document9
    def _rel_str(p: Path) -> str:
        try:
            return p.relative_to(root).as_posix()
        except ValueError:
            return p.as_posix()

    for path in sorted(root.glob(pattern), key=lambda p: _natural_sort_key(_rel_str(p))):
        if path.name.startswith("~$"):
            continue
        if path.resolve() == output_path.resolve():
            continue
        if not include_root_files and path.parent == root:
            continue
        candidates.append(path)
    return candidates


def _extract_summary_document(docx_path: Path) -> Optional[Document]:
    """Generate a document containing only the Summary section, if present."""

    source_doc = Document(docx_path)
    blocks = _iter_block_items(source_doc)
    summary_range = _locate_summary_range(blocks)
    if summary_range is None:
        return None
    start, end = summary_range
    trimmed = _trim_document_to_range_in_place(source_doc, blocks, start, end)
    _label_summary_heading_with_filename(trimmed, docx_path.name)
    return trimmed


def collect_summaries(
    target_root: Path,
    output_path: Path,
    include_root_files: bool = False,
    on_progress: Optional[Callable[[int, int, Optional[Path]], None]] = None,
    should_cancel: Optional[Callable[[], bool]] = None,
    errors: Optional[List[str]] = None,
) -> int:
    """Collect Summary sections from DOCX files under ``target_root``.

    The merge is performed using document-level composition to preserve
    fonts, styles, images, and their sizes.
    """

    docx_files = _gather_docx_files(target_root, include_root_files, output_path)

    total = len(docx_files)
    if on_progress is not None:
        on_progress(0, total, None)

    summary_docs: List[Document] = []
    for idx, docx_path in enumerate(docx_files, start=1):
        if should_cancel is not None and should_cancel():
            raise TimeoutError("Operation was cancelled or timed out.")
        try:
            summary_doc = _extract_summary_document(docx_path)
            if summary_doc is not None:
                summary_docs.append(summary_doc)
        except Exception as exc:
            if errors is not None:
                errors.append(f"{docx_path}: {exc}")
        finally:
            if on_progress is not None:
                on_progress(idx, total, docx_path)

    if not summary_docs:
        doc = Document()
        _clear_document(doc)
        doc.add_paragraph("No Summary sections were found.")
        doc.save(str(output_path))
        return 0

    base = summary_docs[0]
    composer = Composer(base)  # Compose at document level to keep styles/media
    for doc in summary_docs[1:]:
        composer.append(doc)
    composer.save(str(output_path))
    return len(summary_docs)


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    """Parse command-line arguments.

    Positional 'parent' defines the parent folder to scan. If --output is not
    provided, the results are saved as 'collected_summaries.docx' inside the
    parent folder.
    """

    parser = argparse.ArgumentParser(
        description=(
            "Extract 'Summary' sections from .docx files in subfolders of the"
            " specified parent folder and merge them into a single document."
        )
    )
    parser.add_argument(
        "parent",
        help="Parent folder containing subfolders with .docx files",
    )
    parser.add_argument(
        "--output",
        help="Destination DOCX file to create (default: <parent>/collected_summaries.docx)",
    )
    parser.add_argument(
        "--include-root",
        action="store_true",
        help="Also consider .docx files located directly under the parent folder.",
    )
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> None:
    """Entry point for the CLI."""

    args = parse_args(argv)
    target_root = Path(args.parent).expanduser().resolve()
    # Default output: <parent>/collected_summaries.docx when --output is omitted
    if args.output:
        output_path = Path(args.output).expanduser().resolve()
    else:
        output_path = (target_root / "collected_summaries.docx").resolve()

    if not target_root.exists() or not target_root.is_dir():
        raise SystemExit(f"Target directory not found: {target_root}")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    summary_count = collect_summaries(
        target_root=target_root,
        output_path=output_path,
        include_root_files=args.include_root,
    )

    if summary_count == 0:
        print(f"No Summary sections detected under {target_root}.")
    else:
        print(
            f"Collected {summary_count} Summary section{'s' if summary_count != 1 else ''} into {output_path}."
        )


if __name__ == "__main__":
    main(sys.argv[1:])
